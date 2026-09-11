import os
import random
from datetime import datetime, timedelta
from docx import Document
from fpdf import FPDF

# Crear carpeta si no existe
os.makedirs("storage", exist_ok=True)

# Datos semilla realistas para el contexto
empresas = ["Supermercados Mas por Menos", "Alkosto", "Calzatodo", "Cinemark", "Mercado Libre", "SHEIN", "Temu"]
items = ["Mercado general", "Monitor AOC 23.8 LCD", "Protector ECU Yamaha FZ 2.0", "Proyector LED 4K Android", "Sudadera PAVTROS", "Creatina Monohidratada"]
entidades_contratos = ["Forma Atletica", "Autoaprender S.A.S.", "InnovaTech", "DevMasters", "Semillero AZUL"]
objetos_contratos = ["Membresia anual de gimnasio", "Curso de conduccion categorias A2 y C1", "Desarrollo de backend con FastAPI", "Mantenimiento de infraestructura Cloud", "Investigacion en componentes de servicios web"]

nombres = ["Juan Esteban Rojas Melendez", "Carlos Ruiz", "Laura Martinez", "Camila Rojas", "Andres Felipe", "Diana Carolina"]
perfiles = [
    "Estudiante de Tecnologia en Manejo de Sistemas de Informacion en las Unidades Tecnologicas de Santander. Miembro activo del Semillero AZUL.",
    "Ingeniero de Datos con certificaciones de SENATIC y Coursera en IA-native data engineering.",
    "Desarrollador movil especializado en Swift e iOS app development.",
    "Desarrollador Fullstack con experiencia en React, GitHub y Azure foundations."
]

def random_date():
    start = datetime(2024, 1, 1)
    end = datetime(2026, 12, 31)
    return (start + timedelta(days=random.randint(0, (end - start).days))).strftime("%d/%m/%Y")

print("Generando documentos...")

# 1. Generar 10 Facturas en .TXT
for i in range(1, 11):
    empresa = random.choice(empresas)
    item = random.choice(items)
    valor = random.randint(50, 2500) * 1000
    
    contenido = f"FACTURA ELECTRONICA DE VENTA No. FAC-{i*100}\nFecha: {random_date()}\n\nEMISOR:\nRazon Social: {empresa}\nNIT: {random.randint(800,900)}.{random.randint(100,999)}.123-1\n\nDETALLE:\n1x {item} - $ {valor:,} COP\n\nVALOR TOTAL: $ {valor:,} COP"
    
    with open(f"storage/factura_{i:02d}.txt", "w", encoding="utf-8") as f:
        f.write(contenido)

# 2. Generar 10 Contratos en .DOCX (Word)
for i in range(1, 11):
    entidad = random.choice(entidades_contratos)
    objeto = random.choice(objetos_contratos)
    nombre = random.choice(nombres)
    
    doc = Document()
    doc.add_heading(f'CONTRATO DE PRESTACION DE SERVICIOS No. CON-{i*200}', 0)
    doc.add_heading('PARTES INVOLUCRADAS:', level=1)
    doc.add_paragraph(f'Contratante: {entidad}')
    doc.add_paragraph(f'Contratista: {nombre}')
    doc.add_heading('OBJETO DEL CONTRATO:', level=1)
    doc.add_paragraph(f'El presente contrato tiene por objeto: {objeto}.')
    doc.add_heading('VIGENCIA:', level=1)
    doc.add_paragraph(f'Inicio: {random_date()}\nDuracion: 12 meses.')
    doc.add_heading('VALOR:', level=1)
    doc.add_paragraph(f'El valor acordado es de $ {random.randint(1, 15)}.000.000 COP.')
    
    doc.save(f"storage/contrato_{i:02d}.docx")

# 3. Generar 10 Hojas de Vida en .PDF
for i in range(1, 11):
    nombre = random.choice(nombres)
    perfil = random.choice(perfiles)
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    # Función auxiliar para escribir texto compatible en PDF
    def add_text(text, is_bold=False, size=12):
        pdf.set_font("Arial", style='B' if is_bold else '', size=size)
        # Reemplazar caracteres especiales para evitar errores de codificación en FPDF básico
        safe_text = text.encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 8, txt=safe_text)

    add_text("HOJA DE VIDA - PERFIL PROFESIONAL", is_bold=True, size=16)
    pdf.ln(5)
    add_text(f"NOMBRE DEL CANDIDATO: {nombre}", is_bold=True)
    pdf.ln(5)
    add_text("PERFIL:", is_bold=True)
    add_text(perfil)
    pdf.ln(5)
    add_text("EXPERIENCIA Y EDUCACION:", is_bold=True)
    add_text(f"- Mas de {random.randint(1, 5)} anos de experiencia corporativa.")
    add_text("- Liderazgo en proyectos de desarrollo e implementacion.")
    pdf.ln(5)
    add_text("CONTACTO:", is_bold=True)
    add_text(f"Telefono: 31{random.randint(1000000, 9999999)}\nUbicacion: Bucaramanga, Santander.")
    
    pdf.output(f"storage/hojavida_{i:02d}.pdf")

print("¡Éxito! Se crearon 10 TXT, 10 DOCX y 10 PDF en la carpeta 'storage'.")
